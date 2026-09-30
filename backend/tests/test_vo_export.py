import pytest
from fastapi.testclient import TestClient
from main import app
import openpyxl
from io import BytesIO

client = TestClient(app)

def test_export_vo_comparison():
    payload = [
        {
            "start": 1.0,
            "end": 2.5,
            "acceptanceText": "Hello",
            "emissionText": "Hello",
            "isDifferent": False,
            "differenceType": "same"
        },
        {
            "start": 3.0,
            "end": 4.0,
            "acceptanceText": "World",
            "emissionText": "Wurld",
            "isDifferent": True,
            "differenceType": "changed"
        },
        {
            "start": 5.0,
            "end": 6.0,
            "acceptanceText": "Only Acceptance",
            "emissionText": None,
            "isDifferent": True,
            "differenceType": "missing_emission"
        }
    ]

    response = client.post("/api/v1/files/export-vo", json=payload)

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "vo_comparison.xlsx" in response.headers["content-disposition"]

    # Load the workbook from the binary content
    wb = openpyxl.load_workbook(BytesIO(response.content))
    ws = wb.active

    # Check headers
    assert ws.cell(row=1, column=1).value == "Time"
    assert ws.cell(row=1, column=2).value == "Acceptance VO"
    assert ws.cell(row=1, column=3).value == "Emission VO"
    assert ws.cell(row=1, column=4).value == "Difference"

    # Check Row 2 (same)
    assert ws.cell(row=2, column=1).value == "1.0s - 2.5s"
    assert ws.cell(row=2, column=2).value == "Hello"
    assert ws.cell(row=2, column=3).value == "Hello"
    assert ws.cell(row=2, column=4).value == "same"
    # Should not have red formatting on column 4 (differenceType)
    # The default font color is typically None or black/theme

    # Check Row 3 (changed)
    assert ws.cell(row=3, column=1).value == "3.0s - 4.0s"
    assert ws.cell(row=3, column=2).value == "World"
    assert ws.cell(row=3, column=3).value == "Wurld"
    assert ws.cell(row=3, column=4).value == "changed"
    assert ws.cell(row=3, column=4).font.color.rgb == "00FF0000" # Red font
    assert ws.cell(row=3, column=2).fill.start_color.rgb == "00FFCCCC" # Red fill
    assert ws.cell(row=3, column=3).fill.start_color.rgb == "00FFCCCC"

    # Check Row 4 (missing_emission)
    assert ws.cell(row=4, column=2).value == "Only Acceptance"
    assert ws.cell(row=4, column=3).value in ["", None]
    assert ws.cell(row=4, column=4).value == "missing_emission"
    assert ws.cell(row=4, column=3).fill.start_color.rgb == "00FFCCCC" # Red fill on missing side
    assert ws.cell(row=4, column=4).font.color.rgb == "00FF0000"
